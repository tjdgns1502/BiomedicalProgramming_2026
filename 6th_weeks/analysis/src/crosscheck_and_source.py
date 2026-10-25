from pathlib import Path
import json,csv,hashlib,sys,time,warnings,re,urllib.request,xml.etree.ElementTree as ET
import numpy as np
from openpyxl import load_workbook
root=Path(sys.argv[1]).resolve();run=root/'analysis/runs'/sys.argv[2]
def readcsv(p):
 with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
start=time.perf_counter();raw=root/'data/raw/pone.0154420.s001.xlsx';plan=json.loads((root/'analysis/config/plan-v1.0.json').read_text(encoding='utf-8'))
with warnings.catch_warnings(record=True) as caught:
 warnings.simplefilter('always');book=load_workbook(raw,read_only=True,data_only=True)
 sheet=book['Heart_Rate'];source=list(sheet.iter_rows(min_row=2,max_row=23,values_only=True))
rows=readcsv(run/'tables/heart_rate_pairs.csv');rsummary=readcsv(run/'tables/device_summary.csv');checks=[]
def check(name,ok,detail=''):
 checks.append({'check':name,'passed':bool(ok),'detail':detail})
 if not ok:raise AssertionError(name+' '+str(detail))
check('raw_still_identical',sha(raw)==plan['raw_sha256'])
check('rows_88_no_exclusion',len(rows)==88)
devices=['Apple Watch','Fitbit Charge HR','Samsung Gear S','Mio Alpha'];cols=[1,4,7,10]
for name,col in zip(devices,cols):
 vals=[v for v in rows if v['device']==name]
 for orig,z in zip(source,vals):
  for label,origval in [('participant_id',orig[0]),('device_bpm',orig[col]),('reference_bpm',orig[col+1]),('trial',orig[col+2])]:
   check(f'{name}_ID{orig[0]}_{label}',abs(float(z[label])-float(origval))<1e-12)
 x=np.array([v[col+1] for v in source],dtype=float);y=np.array([v[col] for v in source],dtype=float);e=y-x
 beta=np.linalg.lstsq(np.column_stack([np.ones(len(x)),x]),y,rcond=None)[0]
 metrics={'n':len(x),'pearson_r':np.corrcoef(x,y)[0,1],'intercept':beta[0],'slope':beta[1],'bias':e.mean(),'error_sd':e.std(ddof=1),'mae':np.abs(e).mean(),'rmse':np.sqrt(np.mean(e**2)),'mpe':np.mean(100*e/x),'mape':np.mean(100*np.abs(e)/x),'classical_loa_low':e.mean()-1.96*e.std(ddof=1),'classical_loa_high':e.mean()+1.96*e.std(ddof=1)}
 r=next(z for z in rsummary if z['device']==name)
 for key,value in metrics.items():check(name+'_'+key,abs(float(r[key])-float(value))<1e-9,{'R':float(r[key]),'Python':float(value)})
report={'run_id':sys.argv[2],'scope':'same-agent R-vs-Python/OpenPyXL/numpy independent implementation check; not external review','raw_sha256':sha(raw),'checks_passed':len(checks),'elapsed_seconds':time.perf_counter()-start,'reader_warnings':sorted(set(str(x.message) for x in caught)),'decision':'pass','checks':checks}
(run/'python_crosscheck.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
html=(root/'sources/paper.source.txt').read_text(encoding='utf-8')
hrefs=re.findall(r'(?:href|src)=["\']([^"\']*0154420\.t001[^"\']*)',html)
print('TABLE_LINKS',json.dumps(list(dict.fromkeys(hrefs))[:8],ensure_ascii=False))
xmlpath=root/'sources/paper.manuscript.xml'
if not xmlpath.exists():
 url='https://journals.plos.org/plosone/article/file?id=10.1371/journal.pone.0154420&type=manuscript'
 try:
  with urllib.request.urlopen(urllib.request.Request(url,headers={'User-Agent':'Mozilla/5.0'}),timeout=45) as response:data=response.read(2_000_000)
  article=ET.fromstring(data)
  xmlpath.write_bytes(data)
 except Exception as ex:
  (root/'sources/table_acquisition_issue.json').write_text(json.dumps({'url':url,'error':str(ex)},ensure_ascii=False,indent=2),encoding='utf-8')
if xmlpath.exists():
 article=ET.fromstring(xmlpath.read_bytes())
 tables=[]
 for table in article.findall('.//table-wrap'):
  trs=[]
  for tr in table.findall('.//tr'):
   trs.append([' '.join(''.join(c.itertext()).split()) for c in tr if c.tag in ('th','td')])
  rec={'id':table.attrib.get('id'),'label':' '.join(''.join(table.find('label').itertext()).split()) if table.find('label') is not None else '', 'rows':trs,'text':' '.join(''.join(table.itertext()).split())}
  tables.append(rec)
 (root/'sources/published_tables.json').write_text(json.dumps(tables,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
 print('PUBLISHED_TABLES',json.dumps(tables,ensure_ascii=False)[:10500])
print('CROSSCHECK',json.dumps({k:v for k,v in report.items() if k!='checks'},ensure_ascii=False))
print('SECONDARY_RHO',[(x['device'],x['spearman_rho'],x['bias'],x['classical_loa_low'],x['classical_loa_high'],x['ba_bias_slope_p']) for x in rsummary])
